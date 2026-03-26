from __future__ import annotations

import re
from pathlib import Path

from docx import Document


INPUT_DOC = Path("/Users/emmanuel/Desktop/apps/gerts/output/doc/GERTS_thesis_reworked_final_v7.docx")
OUTPUT_DOC = Path("/Users/emmanuel/Desktop/apps/gerts/output/doc/GERTS_thesis_reworked_final_v8.docx")


ORDERED_OLD_NUMBERS = [
    "2.1",
    "2.7",
    "2.2",
    "2.8",
    "2.9",
    "2.3",
    "2.10",
    "2.4",
    "2.11",
    "2.12",
    "2.5",
    "2.13",
    "2.6",
    "2.14",
    "2.15",
    "2.16",
    "2.17",
]


def replace_table_refs(text: str, mapping: dict[str, str]) -> str:
    placeholders = {}
    for idx, (old, new) in enumerate(mapping.items(), start=1):
        placeholder = f"__TABREF_{idx}__"
        placeholders[placeholder] = f"Table {new}"
        text = re.sub(rf"\bTable {re.escape(old)}\b", placeholder, text)
    for placeholder, replacement in placeholders.items():
        text = text.replace(placeholder, replacement)
    return text


def main() -> None:
    document = Document(INPUT_DOC)
    mapping = {old: f"2.{idx}" for idx, old in enumerate(ORDERED_OLD_NUMBERS, start=1)}

    for paragraph in document.paragraphs:
        text = paragraph.text
        if "Table 2." not in text:
            continue
        paragraph.text = replace_table_refs(text, mapping)

    OUTPUT_DOC.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT_DOC)
    print(OUTPUT_DOC)


if __name__ == "__main__":
    main()
